import argparse
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from update_augmented_cost_estimation import (
    HEADERS,
    build_values,
    read_baseline,
    same_configuration,
    upsert_result,
)


def arguments(**overrides):
    values = {
        "test_db": "accidents",
        "cardinality_type": "est",
        "time_stamp": "20260820_120000",
        "epochs": 100,
        "test_augment": "True",
        "augment_pooling": "hybrid",
        "augment_refinement": "gated_residual",
        "augment_coarse_layers": 1,
        "augment_include_inv": "False",
        "augment_refine_ret": "False",
        "lambda_struct": 0.0,
        "activation": "LeakyReLU",
        "augment_mq_queries": 8,
        "augment_seq_regions": "False",
        "augment_cfg_coarse_edges": "False",
    }
    values.update(overrides)
    return argparse.Namespace(**values)


class UpdateAugmentedCostEstimationTest(unittest.TestCase):
    def test_legacy_config_defaults_test_augment_to_true(self):
        current = {
            "test_db": "accidents",
            "cardinality_type": "est",
            "epochs": 100,
            "test-augment": "True",
            "augment-pooling": "hybrid",
            "augment-refinement": "gated_residual",
            "augment-coarse-layers": 1,
            "augment-include-inv": "False",
            "augment-refine-ret": "False",
            "lambda-struct": 0.0,
            "activation": "LeakyReLU",
        }
        legacy = {**current, "test-augment": None}

        self.assertTrue(same_configuration(legacy, current))
        self.assertFalse(same_configuration(legacy, {**current, "test-augment": "False"}))

    def test_read_baseline_matches_database_and_cardinality(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "baseline.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Baseline"
            worksheet.append([
                "test_db", "cardinality_type",
                "pull_up_q50_mean", "pull_up_q95_mean", "pull_up_q99_mean",
                "push_down_q50_mean", "push_down_q95_mean", "push_down_q99_mean",
                "time_stamp",
            ])
            worksheet.append(["accidents", "act", 20.0, 30.0, 40.0, 50.0, 60.0, 70.0, "old"])
            worksheet.append(["accidents", "est", 4.0, 7.0, 10.0, 5.0, 8.0, 11.0, "new"])
            workbook.save(path)
            workbook.close()

            actual = read_baseline(path, "ACCIDENTS", "EST")

            self.assertEqual(actual["pullup"], {"q50": 4.0, "q95": 7.0, "q99": 10.0})
            self.assertEqual(actual["pushdown"], {"q50": 5.0, "q95": 8.0, "q99": 11.0})

    def test_build_values_computes_baseline_minus_augmented(self):
        summary = {
            "workloads": {
                "workload_pullup_est": {
                    "q50": (2.5, 0.1, 2), "q95": (5.0, 0.2, 2), "q99": (8.5, 0.3, 2),
                },
                "workload_pushdown_est": {
                    "q50": (3.0, 0.1, 2), "q95": (6.0, 0.2, 2), "q99": (9.0, 0.3, 2),
                },
            }
        }

        actual = build_values(
            arguments(),
            summary,
            {
                "pullup": {"q50": 4.0, "q95": 7.0, "q99": 10.0},
                "pushdown": {"q50": 5.0, "q95": 8.0, "q99": 11.0},
            },
        )

        self.assertEqual(actual["diff_pull_up_q50_mean"], 1.5)
        self.assertEqual(actual["diff_pull_up_q95_mean"], 2.0)
        self.assertEqual(actual["diff_pull_up_q99_mean"], 1.5)
        self.assertEqual(actual["diff_push_down_q50_mean"], 2.0)

    def test_upsert_keeps_distinct_configs_and_replaces_matching_config(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "augmented.xlsx"
            summary = {
                "workloads": {
                    "workload_pullup_est": {
                        "q50": (2.5, 0.0, 1), "q95": (5.0, 0.0, 1), "q99": (8.5, 0.0, 1),
                    },
                    "workload_pushdown_est": {
                        "q50": (3.0, 0.0, 1), "q95": (6.0, 0.0, 1), "q99": (9.0, 0.0, 1),
                    },
                }
            }
            baseline = {
                "pullup": {"q50": 4.0, "q95": 7.0, "q99": 10.0},
                "pushdown": {"q50": 5.0, "q95": 8.0, "q99": 11.0},
            }
            first = build_values(arguments(), summary, baseline)
            replacement = build_values(arguments(time_stamp="20260820_130000"), summary, baseline)
            distinct_epochs = build_values(arguments(epochs=200), summary, baseline)
            distinct = build_values(arguments(augment_pooling="max"), summary, baseline)
            no_test_augment = build_values(arguments(test_augment="False"), summary, baseline)

            upsert_result(str(path), first)
            upsert_result(str(path), replacement)
            upsert_result(str(path), distinct_epochs)
            upsert_result(str(path), distinct)
            upsert_result(str(path), no_test_augment)

            workbook = load_workbook(path, read_only=True, data_only=True)
            worksheet = workbook["Augmented"]
            rows = list(worksheet.iter_rows(values_only=True))
            workbook.close()
            self.assertEqual(list(rows[0]), HEADERS)
            self.assertEqual(len(rows), 5)
            self.assertEqual(rows[1][HEADERS.index("time_stamp")], "20260820_130000")
            self.assertEqual(rows[1][HEADERS.index("epochs")], 100)
            self.assertEqual(rows[2][HEADERS.index("epochs")], 200)
            self.assertEqual(rows[3][HEADERS.index("augment-pooling")], "max")
            self.assertEqual(rows[4][HEADERS.index("test-augment")], "False")

    def test_legacy_blank_activation_is_kept_not_overwritten(self):
        current = {
            "test_db": "accidents",
            "cardinality_type": "est",
            "epochs": 100,
            "test-augment": "True",
            "augment-pooling": "hybrid",
            "augment-refinement": "gated_residual",
            "augment-coarse-layers": 1,
            "augment-include-inv": "False",
            "augment-refine-ret": "False",
            "lambda-struct": 0.0,
            "activation": "LeakyReLU",
        }
        legacy = {**current, "activation": None}

        self.assertFalse(same_configuration(legacy, current))
        self.assertTrue(same_configuration(current, {**current, "activation": "leakyrelu"}))

    def test_upsert_keeps_one_row_per_activation(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "augmented.xlsx"
            summary = {
                "workloads": {
                    "workload_pullup_est": {
                        "q50": (2.5, 0.0, 1), "q95": (5.0, 0.0, 1), "q99": (8.5, 0.0, 1),
                    },
                    "workload_pushdown_est": {
                        "q50": (3.0, 0.0, 1), "q95": (6.0, 0.0, 1), "q99": (9.0, 0.0, 1),
                    },
                }
            }
            baseline = {
                "pullup": {"q50": 4.0, "q95": 7.0, "q99": 10.0},
                "pushdown": {"q50": 5.0, "q95": 8.0, "q99": 11.0},
            }
            for activation in ("LeakyReLU", "ReLU", "SELU", "CELU"):
                upsert_result(str(path), build_values(arguments(activation=activation), summary, baseline))

            workbook = load_workbook(path, read_only=True, data_only=True)
            rows = list(workbook["Augmented"].iter_rows(values_only=True))
            workbook.close()
            self.assertEqual(
                [row[HEADERS.index("activation")] for row in rows[1:]],
                ["LeakyReLU", "ReLU", "SELU", "CELU"],
            )


    def test_mq_queries_blank_for_other_poolings(self):
        # run_code.sh always passes AUGMENT_MQ_QUERIES; only MQ rows should record it.
        values = build_values(arguments(augment_pooling="attention"), MQ_SUMMARY, MQ_BASELINE)
        self.assertIsNone(values["augment-mq-queries"])

        legacy = {header: values.get(header) for header in HEADERS if header != "augment-mq-queries"}
        self.assertTrue(same_configuration(legacy, values))

        mq = build_values(arguments(augment_pooling="multi_query_attention"), MQ_SUMMARY, MQ_BASELINE)
        self.assertEqual(mq["augment-mq-queries"], 8)

    def test_upsert_keeps_one_row_per_mq_query_count(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "augmented.xlsx"
            upsert_result(str(path), build_values(arguments(augment_pooling="attention"), MQ_SUMMARY, MQ_BASELINE))
            for num_queries in (4, 8, 16):
                upsert_result(str(path), build_values(
                    arguments(augment_pooling="multi_query_attention", augment_mq_queries=num_queries),
                    MQ_SUMMARY, MQ_BASELINE))
            # A rerun of M=8 replaces its own row; it does not touch M=4/16 or the attention row.
            upsert_result(str(path), build_values(
                arguments(augment_pooling="multi_query_attention", augment_mq_queries=8,
                          time_stamp="20260820_130000"),
                MQ_SUMMARY, MQ_BASELINE))

            workbook = load_workbook(path, read_only=True, data_only=True)
            rows = list(workbook["Augmented"].iter_rows(values_only=True))
            workbook.close()
            self.assertEqual(
                [(row[HEADERS.index("augment-pooling")], row[HEADERS.index("augment-mq-queries")])
                 for row in rows[1:]],
                [("attention", None), ("multi_query_attention", 4),
                 ("multi_query_attention", 8), ("multi_query_attention", 16)],
            )
            self.assertEqual(rows[3][HEADERS.index("time_stamp")], "20260820_130000")

    def test_legacy_blank_supernode_flags_are_kept_not_overwritten(self):
        summary = {"workloads": {
            "workload_pullup_est": {"q50": (2.5, 0.0, 1), "q95": (5.0, 0.0, 1), "q99": (8.5, 0.0, 1)},
            "workload_pushdown_est": {"q50": (3.0, 0.0, 1), "q95": (6.0, 0.0, 1), "q99": (9.0, 0.0, 1)},
        }}
        baseline = {"pullup": {"q50": 4.0, "q95": 7.0, "q99": 10.0},
                    "pushdown": {"q50": 5.0, "q95": 8.0, "q99": 11.0}}
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "augmented.xlsx"
            legacy = build_values(arguments(augment_seq_regions=None, augment_cfg_coarse_edges=None),
                                  summary, baseline)
            upsert_result(str(path), legacy)
            for seq, cfg in (("True", "True"), ("False", "False"), ("True", "True")):
                upsert_result(str(path), build_values(
                    arguments(augment_seq_regions=seq, augment_cfg_coarse_edges=cfg), summary, baseline))

            workbook = load_workbook(path, read_only=True, data_only=True)
            rows = list(workbook["Augmented"].iter_rows(values_only=True))
            workbook.close()
            flags = [(row[HEADERS.index("augment-seq-regions")], row[HEADERS.index("augment-cfg-coarse-edges")])
                     for row in rows[1:]]
            self.assertEqual(flags, [(None, None), ("True", "True"), ("False", "False")])



MQ_SUMMARY = {
    "workloads": {
        "workload_pullup_est": {"q50": (2.5, 0.0, 1), "q95": (5.0, 0.0, 1), "q99": (8.5, 0.0, 1)},
        "workload_pushdown_est": {"q50": (3.0, 0.0, 1), "q95": (6.0, 0.0, 1), "q99": (9.0, 0.0, 1)},
    }
}
MQ_BASELINE = {
    "pullup": {"q50": 4.0, "q95": 7.0, "q99": 10.0},
    "pushdown": {"q50": 5.0, "q95": 8.0, "q99": 11.0},
}


if __name__ == "__main__":
    unittest.main()
